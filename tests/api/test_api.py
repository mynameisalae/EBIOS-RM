"""L'interface web : ce que le serveur lit, produit et refuse (conception §20).

A scripted runner stands in for the model; the mission database is a real SQLite file.
"""

from __future__ import annotations

import io
import threading
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient

from ebios_rm.api.app import create_app
from ebios_rm.api.views import atelier_metrics
from ebios_rm.repositories.mission_repository import MissionRepository, connect
from ebios_rm.reporting.audit_annex import build_audit_annex, collect_audit_annex_data
from ebios_rm.reporting.report_agent import SectionDraft, Statement

W4 = {"scenarios": [
    {"id": "SO-01", "statut": "confirme", "resume": "Via le serveur web",
     "anomalies": [{"code": "revision_forte", "message": "V4 -> V2", "bloquante": False}],
     "attack_path": [{"technique_id": "T1190", "technique_name": "Exploit Public-Facing Application"},
                     {"technique_id": "T1041", "technique_name": "Exfiltration Over C2 Channel"}]},
    {"id": "SO-02", "statut": "confirme", "resume": "Via un compte",
     "attack_path": [{"technique_id": "T1078", "technique_name": "Valid Accounts"},
                     {"technique_id": "T1041", "technique_name": "Exfiltration Over C2 Channel"}]},
]}
W5 = {"risques": [
    {"id": "R1", "niveau_risque": "Critique", "niveau_risque_residuel": "Élevé", "vraisemblance_residuelle": "V2"},
    {"id": "R2", "niveau_risque": "Critique", "niveau_risque_residuel": "Élevé", "vraisemblance_residuelle": "V2"},
], "mesures": [{"id": "M-01"}, {"id": "M-02"}, {"id": "M-03"}],
    "cadre_suivi": {"indicateurs": [{"id": "IND-01"}, {"id": "IND-02"}]}}


class Scripted:
    def __init__(self, hold: threading.Event | None = None) -> None:
        self.hold = hold

    def write_section(self, spec, state):
        if self.hold is not None:
            self.hold.wait(5)
        if spec.key == "synthese":
            return SectionDraft(statements=[Statement(text="Deux risques sont retenus : R1 et R2.", sources=["w5.risques"])])
        if spec.key == "appreciations":
            return SectionDraft(statements=[Statement(
                text="L'auditeur a validé la correction.",
                sources=["auditeur.appreciations.0.appreciation"])])
        return SectionDraft(statements=[])

    def revise_section(self, spec, state, refused):
        return SectionDraft(statements=[])


@pytest.fixture
def db(tmp_path):
    path = str(tmp_path / "mission.db")
    repo = MissionRepository(connect(path))
    mid = repo.create_mission("mondial_test", [])
    repo.save_output(mid, 0, {"organisation_nom": "Mondial-Tickets", "secteur_activite": "Billetterie",
                              "applicable_frameworks": []})
    for n, output in ((4, W4), (5, W5)):
        version = repo.save_output(mid, n, output)
        repo.set_version_status(mid, n, version, "approved")
    repo.save_output(mid, 1, {"biens_essentiels": [{"id": "BE-1"}]})          # exists, not approved
    repo.log_decision(mid, stage="w4", action="approved", justification="ok")
    return path, mid


@pytest.fixture
def client(db, tmp_path):
    path, _ = db
    return TestClient(create_app(path, reports_dir=tmp_path / "reports", runner_factory=lambda progress: Scripted()))


def _wait(client, job_id, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] != "running":
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


# --- reading -----------------------------------------------------------------------------

def test_the_missions_list_shows_each_atelier_state(client, db):
    _, mid = db

    [mission] = client.get("/api/missions").json()

    assert mission["id"] == mid and mission["organisation"] == "Mondial-Tickets"
    assert {a["n"]: a["state"] for a in mission["ateliers"]} == {
        1: "in_progress", 2: "absent", 3: "absent", 4: "approved", 5: "approved"}
    assert mission["approved"] == 2


def test_the_overview_gives_real_figures_per_atelier(client, db):
    _, mid = db

    overview = client.get(f"/api/missions/{mid}").json()

    by_n = {a["n"]: a for a in overview["ateliers"]}
    assert by_n[4]["metrics"] == [{"label": "Modes opératoires", "value": 2}, {"label": "Étapes d'attaque", "value": 4},
                                  {"label": "Techniques ATT&CK", "value": 3}]   # T1041 is cited twice, counted once
    assert by_n[4]["headline"] == "2 modes opératoires, 3 techniques ATT&CK distinctes"
    assert by_n[5]["headline"] == "2 risques ramenés de Critique à Élevé"
    assert by_n[2]["state"] == "absent" and by_n[2]["metrics"] == []
    assert overview["organisation"]["nom"] == "Mondial-Tickets"
    assert overview["decisions"] == 1 and overview["can_write_report"] is True


def test_an_atelier_detail_carries_its_output_and_its_versions(client, db):
    _, mid = db

    detail = client.get(f"/api/missions/{mid}/ateliers/4").json()

    assert detail["output"]["scenarios"][0]["id"] == "SO-01"
    assert [(v["number"], v["status"]) for v in detail["versions"]] == [(1, "approved")]
    assert detail["question"].startswith("Comment, techniquement")


def test_unknown_things_are_a_404(client, db):
    _, mid = db
    assert client.get("/api/missions/inconnue").status_code == 404
    assert client.get(f"/api/missions/{mid}/ateliers/9").status_code == 404
    assert client.get(f"/api/missions/{mid}/ateliers/0").status_code == 404
    assert client.get("/api/missions/inconnue/annexe").status_code == 404
    assert client.get("/api/jobs/inconnu").status_code == 404


@pytest.mark.parametrize("n, output, expected_labels", [
    (1, {"biens_essentiels": [1, 2], "biens_supports": [1], "evenements_redoutes": [1], "baseline_gaps_full": []},
     ["Biens essentiels", "Biens supports", "Événements redoutés", "Écarts du socle"]),
    (2, {"sources_risque": [1], "objectifs_vises": [1], "couples": [1]},
     ["Sources de risque", "Objectifs visés", "Couples retenus"]),
    (3, {"scenarios": [1, 2, 3]}, ["Scénarios stratégiques"]),
])
def test_atelier_figures_are_counted_from_the_stored_output(n, output, expected_labels):
    _, metrics = atelier_metrics(n, output)
    assert [m["label"] for m in metrics] == expected_labels


def test_the_project_page_has_the_real_frameworks(client):
    projet = client.get("/api/projet").json()

    assert {f["id"] for f in projet["frameworks"]} >= {"ISO27001", "RGPD", "NIST", "ANSSI_hygiene"}
    assert projet["controls"] == sum(f["controls"] for f in projet["frameworks"]) > 0
    assert [a["n"] for a in projet["ateliers"]] == [1, 2, 3, 4, 5]


# --- the two documents -------------------------------------------------------------------

def test_the_audit_annex_is_downloadable_and_readable(client, db):
    _, mid = db

    response = client.get(f"/api/missions/{mid}/annexe")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f"annexe_audit_{mid[:8]}.docx" in response.headers["content-disposition"]
    assert "mondial_test" in [p.text for p in Document(io.BytesIO(response.content)).paragraphs]


def test_a_report_is_written_in_the_background_then_downloaded(client, db):
    _, mid = db

    started = client.post(f"/api/missions/{mid}/rapport").json()
    job = _wait(client, started["job_id"])

    assert job["status"] == "done" and job["summary"]["statements"] == 1
    file = client.get(f"/api/jobs/{started['job_id']}/fichier")
    assert file.status_code == 200
    assert "Deux risques sont retenus : R1 et R2. [1]" in [p.text for p in Document(io.BytesIO(file.content)).paragraphs]


def test_nothing_to_report_without_an_approved_atelier(tmp_path):
    path = str(tmp_path / "vide.db")
    repo = MissionRepository(connect(path))
    mid = repo.create_mission("vide", [])
    repo.save_output(mid, 1, {})
    client = TestClient(create_app(path, reports_dir=tmp_path, runner_factory=lambda p: Scripted()))

    assert client.post(f"/api/missions/{mid}/rapport").status_code == 409


def test_a_filled_annex_adds_the_auditors_appraisals(client, db):
    path, mid = db
    doc = build_audit_annex(collect_audit_annex_data(MissionRepository(connect(path)), mid, framework_versions={}))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Origine")
    table.rows[1].cells[3].text = "Correction vérifiée."
    buffer = io.BytesIO()
    doc.save(buffer)

    started = client.post(f"/api/missions/{mid}/rapport", files={"annexe": ("annexe.docx", buffer.getvalue())}).json()
    job = _wait(client, started["job_id"])

    assert started["appraisals"] == 1
    assert job["status"] == "done" and job["summary"]["statements"] == 2   # the synthesis + the appraisal


def test_an_annex_of_another_mission_is_refused(client, db, tmp_path):
    other = MissionRepository(connect(str(tmp_path / "autre.db")))
    other_id = other.create_mission("autre", [])
    buffer = io.BytesIO()
    build_audit_annex(collect_audit_annex_data(other, other_id, framework_versions={})).save(buffer)
    _, mid = db

    response = client.post(f"/api/missions/{mid}/rapport", files={"annexe": ("a.docx", buffer.getvalue())})

    assert response.status_code == 400 and "autre mission" in response.json()["detail"]


def test_a_file_that_is_not_word_is_refused(client, db):
    _, mid = db

    response = client.post(f"/api/missions/{mid}/rapport", files={"annexe": ("a.docx", b"pas un docx")})

    assert response.status_code == 400


def test_only_one_report_is_written_at_a_time(db, tmp_path):
    path, mid = db
    hold = threading.Event()
    client = TestClient(create_app(path, reports_dir=tmp_path, runner_factory=lambda p: Scripted(hold)))

    first = client.post(f"/api/missions/{mid}/rapport")
    second = client.post(f"/api/missions/{mid}/rapport")
    hold.set()

    assert first.status_code == 200 and second.status_code == 409
    assert _wait(client, first.json()["job_id"])["status"] == "done"
