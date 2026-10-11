"""Fiche de test — la boucle de l'auditeur (conception §20.2).

Annex out with blank cells, the auditor fills some in Word, code reads back the filled
ones, the model writes them up from their own slice of the mission, the gate checks.
Everything but the model is exercised for real, on actual .docx files.
"""

from __future__ import annotations

import pytest
from docx import Document

from ebios_rm.repositories.mission_repository import MissionRepository, connect
from ebios_rm.reporting.appraisals import Appraisal, AnnexReadError, read_appraisals
from ebios_rm.reporting.audit_annex import build_audit_annex, collect_audit_annex_data
from ebios_rm.reporting.mission_report import build_mission_report, build_snapshot, write_mission_report
from ebios_rm.reporting.report_agent import SectionDraft, Statement

FIXED_DATE = "2026-10-10T00:00:00+00:00"


@pytest.fixture
def repo() -> MissionRepository:
    return MissionRepository(connect(":memory:"))


@pytest.fixture
def mission(repo):
    mid = repo.create_mission("Clinique Test", [])
    repo.save_output(mid, 0, {"organisation_nom": "Clinique du Parc", "secteur_activite": "Santé",
                              "applicable_frameworks": []})
    version = repo.save_output(mid, 4, {"scenarios": [
        {"id": "SO-02", "statut": "confirme", "resume": "Accès par compte partagé",
         "anomalies": [{"code": "tactique_incoherente", "message": "T1078 mal classée", "bloquante": False}]},
        {"id": "SO-05", "statut": "confirme", "resume": "Via un tiers",
         "anomalies": [{"code": "vraisemblance_divergente", "message": "V2 annoncée, V3 calculée", "bloquante": False}]},
    ]})
    repo.set_version_status(mid, 4, version, "approved")
    repo.log_decision(mid, stage="w2", action="rejected", justification="Couples mal notés.")
    return mid


def _filled_annex(repo, mid, tmp_path, answers: dict[int, str]):
    """Generate the annex, then write ``answers`` (row index -> text) in the appraisal column, like the auditor."""
    doc = build_audit_annex(collect_audit_annex_data(repo, mid, framework_versions={}, generated_at=FIXED_DATE))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Origine")
    for row_index, text in answers.items():
        table.rows[row_index].cells[3].text = text
    path = tmp_path / "annexe_remplie.docx"
    doc.save(str(path))
    return path


# --- reading back ----------------------------------------------------------------------

def test_only_the_cells_the_auditor_filled_are_read(repo, mission, tmp_path):
    path = _filled_annex(repo, mission, tmp_path, {2: "Vérifié : pas d'impact sur SO-05."})

    reading = read_appraisals(path)

    assert reading.mission_id == mission
    assert reading.appraisals == (
        Appraisal("Atelier 4 · SO-05", "vraisemblance_divergente", "V2 annoncée, V3 calculée",
                  "Vérifié : pas d'impact sur SO-05."),
    )


def test_every_filled_cell_is_read_in_order_and_trimmed(repo, mission, tmp_path):
    path = _filled_annex(repo, mission, tmp_path, {1: "  Corrigé à la main.  ", 3: "Acceptée par la direction."})

    appraisals = read_appraisals(path).appraisals

    assert [(a.origin, a.text) for a in appraisals] == [
        ("Atelier 4 · SO-02", "Corrigé à la main."),
        ("Journal des décisions · w2", "Acceptée par la direction."),
    ]


def test_a_cell_with_only_spaces_is_not_an_appraisal(repo, mission, tmp_path):
    path = _filled_annex(repo, mission, tmp_path, {1: "   ", 2: "\n"})

    assert read_appraisals(path).appraisals == ()


def test_an_annex_without_abnormal_events_has_nothing_to_appraise(repo, tmp_path):
    mid = repo.create_mission("Mission calme", [])
    doc = build_audit_annex(collect_audit_annex_data(repo, mid, framework_versions={}, generated_at=FIXED_DATE))
    path = tmp_path / "calme.docx"
    doc.save(str(path))

    reading = read_appraisals(path)

    assert reading.mission_id == mid and reading.appraisals == ()


def test_a_file_that_is_not_an_annex_is_refused(tmp_path):
    path = tmp_path / "autre.docx"
    Document().save(str(path))

    with pytest.raises(AnnexReadError):
        read_appraisals(path)


# --- writing them into the report ------------------------------------------------------

class AppraisalRunner:
    def __init__(self, statements: list[Statement], revisions: list[Statement] | None = None) -> None:
        self.statements, self.revisions = statements, revisions or []
        self.received: dict[str, dict] = {}

    def write_section(self, spec, state):
        self.received[spec.key] = state
        return SectionDraft(statements=self.statements if spec.key == "appreciations" else [])

    def revise_section(self, spec, state, refused):
        return SectionDraft(statements=self.revisions)


def _read(repo, mission, tmp_path, answers):
    return read_appraisals(_filled_annex(repo, mission, tmp_path, answers)).appraisals


def test_the_model_gets_the_appraisals_and_only_the_atelier_they_concern(repo, mission, tmp_path):
    appraisals = _read(repo, mission, tmp_path, {1: "Corrigé à la main."})
    runner = AppraisalRunner([])

    write_mission_report(build_snapshot(repo, mission), runner, appraisals=appraisals, generated_at=FIXED_DATE)

    received = runner.received["appreciations"]
    assert set(received) == {"organisation", "echelles", "w4", "auditeur"}
    assert received["auditeur"]["appreciations"] == [{
        "origine": "Atelier 4 · SO-02", "code": "tactique_incoherente",
        "message": "T1078 mal classée", "appreciation": "Corrigé à la main."}]


def test_no_appraisal_means_no_model_call_for_that_section(repo, mission):
    runner = AppraisalRunner([])

    report = write_mission_report(build_snapshot(repo, mission), runner, generated_at=FIXED_DATE)

    assert "appreciations" not in runner.received
    assert "appreciations" not in {s.key for s in report.sections}


def test_an_appraisal_statement_must_cite_the_auditors_own_words(repo, mission, tmp_path):
    appraisals = _read(repo, mission, tmp_path, {1: "Corrigé à la main."})
    runner = AppraisalRunner([
        Statement(text="L'auditeur a corrigé SO-02 à la main.",
                  sources=["auditeur.appreciations.0.appreciation", "w4.scenarios.0.id"]),
        Statement(text="SO-02 est confirmé.", sources=["w4.scenarios.0.id"]),   # valid fields, but not the auditor
    ])

    report = write_mission_report(build_snapshot(repo, mission), runner, appraisals=appraisals, generated_at=FIXED_DATE)

    section = next(s for s in report.sections if s.key == "appreciations")
    assert [s.text for s in section.statements] == ["L'auditeur a corrigé SO-02 à la main."]
    assert [r.reason for r in report.rejected] == [
        "abandonnée après refus — ne cite pas l'appréciation de l'auditeur"]


def test_an_appraisal_statement_still_goes_through_the_ordinary_gate(repo, mission, tmp_path):
    appraisals = _read(repo, mission, tmp_path, {1: "Corrigé à la main."})
    runner = AppraisalRunner([
        Statement(text="L'auditeur a corrigé M-99.", sources=["auditeur.appreciations.0.appreciation"]),
    ])

    report = write_mission_report(build_snapshot(repo, mission), runner, appraisals=appraisals, generated_at=FIXED_DATE)

    assert [r.reason for r in report.rejected] == [
        "abandonnée après refus — « M-99 » ne figure pas dans les champs cités"]


def test_the_appraisal_section_reaches_the_document(repo, mission, tmp_path):
    appraisals = _read(repo, mission, tmp_path, {1: "Corrigé à la main."})
    runner = AppraisalRunner([Statement(text="L'auditeur a corrigé l'événement à la main.",
                                        sources=["auditeur.appreciations.0.appreciation"])])
    report = write_mission_report(build_snapshot(repo, mission), runner, appraisals=appraisals, generated_at=FIXED_DATE)

    lines = [p.text for p in build_mission_report(report).paragraphs]

    assert "1. Appréciations de l'auditeur" in lines
    assert "L'auditeur a corrigé l'événement à la main. [1]" in lines
    assert "[1]  auditeur.appreciations.0.appreciation" in lines
