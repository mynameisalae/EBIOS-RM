"""Fiche de test — Annexe d'audit, fully deterministic (conception §20.3).

No model call anywhere: the annex is data in, Word document out, so every
assertion below reads the produced document back.
"""

from __future__ import annotations

import pytest
from docx import Document

from ebios_rm.plugins.registry import discover_frameworks
from ebios_rm.repositories.mission_repository import MissionRepository, connect
from ebios_rm.reporting.house_style import french_month_year
from ebios_rm.reporting.audit_annex import (
    UNKNOWN_VERSION,
    build_audit_annex,
    collect_audit_annex_data,
    write_audit_annex,
)

FIXED_DATE = "2026-10-10T00:00:00+00:00"
CONCEPTION_FRAMEWORKS = {"ISO27001": "2022", "ANSSI_hygiene": "2024", "RGPD": "2016/679", "NIST": "CSF 2.0"}


@pytest.fixture
def repo() -> MissionRepository:
    return MissionRepository(connect(":memory:"))


def _table(doc, first_header: str, columns: int | None = None):
    for table in doc.tables:
        if table.rows[0].cells[0].text == first_header and columns in (None, len(table.columns)):
            return table
    raise AssertionError(f"no table starting with {first_header!r}")


def _rows(doc, first_header: str, columns: int | None = None) -> list[list[str]]:
    """Rows (header excluded) of the table whose first column header is ``first_header``."""
    return [[c.text for c in row.cells] for row in _table(doc, first_header, columns).rows[1:]]


def _annex(repo, mission_id, **kwargs):
    kwargs.setdefault("framework_versions", {})
    kwargs.setdefault("generated_at", FIXED_DATE)
    return build_audit_annex(collect_audit_annex_data(repo, mission_id, **kwargs))


# --- the conception's own test sheet (§20.3) ---------------------------------

def test_versions_labeled_and_framework_versions_rendered(repo):
    mid = repo.create_mission("Clinique Test", ["RGPD", "NIST"])
    repo.save_output(mid, 3, {"scenarios": []})
    v2 = repo.save_output(mid, 3, {"scenarios": []})
    repo.set_version_status(mid, 3, v2, "approved")
    repo.set_attck_version(mid, "v15.1")

    doc = _annex(repo, mid, framework_versions=CONCEPTION_FRAMEWORKS)

    atelier3 = [r for r in _rows(doc, "Atelier") if r[0].startswith("Atelier 3")]
    assert [(r[1], r[2]) for r in atelier3] == [("1", "superseded"), ("2", "approved")]

    versions = dict(_rows(doc, "Référentiel", columns=2))
    assert versions["RGPD"] == "2016/679"
    assert versions["NIST"] == "CSF 2.0"
    assert versions == CONCEPTION_FRAMEWORKS

    header = _rows(doc, "Rubrique")
    assert ["Référentiel ATT&CK utilisé", "v15.1"] in header


def test_every_abnormal_event_has_a_blank_auditor_field(repo):
    mid = repo.create_mission("Clinique Test", [])
    repo.save_output(mid, 4, {"scenarios": [
        {"id": "SO-02", "anomalies": [{"code": "tactique_incoherente", "message": "m1", "bloquante": False}]},
        {"id": "SO-05", "anomalies": [{"code": "vraisemblance_divergente", "message": "m2", "bloquante": False}]},
    ]})
    repo.log_decision(mid, stage="w4", action="rejected", justification="Trop de doublons.")

    doc = _annex(repo, mid)

    events = _rows(doc, "Origine")
    assert len(events) == 3
    assert all(row[3] == "" for row in events)
    assert _table(doc, "Origine").rows[0].cells[3].text == "Appréciation de l'auditeur"


# --- what counts as an abnormal event -----------------------------------------

def test_anomalies_are_attributed_to_their_scenario(repo):
    mid = repo.create_mission("Clinique Test", [])
    repo.save_output(mid, 4, {"scenarios": [
        {"id": "SO-02", "anomalies": [{"code": "tactique_incoherente", "message": "m1", "bloquante": False}]},
    ]})

    events = _rows(_annex(repo, mid), "Origine")

    assert events == [["Atelier 4 · SO-02", "tactique_incoherente", "m1", ""]]


def test_quality_checks_not_ok_are_events_and_ok_ones_are_not(repo):
    mid = repo.create_mission("Clinique Test", [])
    repo.save_output(mid, 4, {"scenarios": [], "quality_report": {"checks": [
        {"controle": "Couverture", "statut": "ok", "message": "tout va bien"},
        {"controle": "Égalité", "statut": "avertissement", "message": "à trancher"},
    ]}})

    events = _rows(_annex(repo, mid), "Origine")

    assert events == [["Atelier 4 · contrôle qualité (avertissement)", "Égalité", "à trancher", ""]]


def test_only_non_routine_decisions_are_abnormal_events(repo):
    mid = repo.create_mission("Clinique Test", [])
    repo.log_decision(mid, stage="w1", action="approved", justification="Rien à signaler.")
    repo.log_decision(mid, stage="w2", action="rejected", justification="Couples mal notés.")

    doc = _annex(repo, mid)

    assert _rows(doc, "Origine") == [["Journal des décisions · w2", "rejected", "Couples mal notés.", ""]]
    log = _rows(doc, "Horodatage")
    assert [(r[1], r[3], r[4]) for r in log] == [
        ("w1", "approved", "Rien à signaler."), ("w2", "rejected", "Couples mal notés."),
    ]


def test_events_come_from_the_current_version_only(repo):
    mid = repo.create_mission("Clinique Test", [])
    repo.save_output(mid, 4, {"scenarios": [{"id": "SO-01", "anomalies": [{"code": "ancienne", "message": "m"}]}]})
    repo.save_output(mid, 4, {"scenarios": [{"id": "SO-01", "anomalies": []}]})

    doc = _annex(repo, mid)

    assert not any(t.rows[0].cells[0].text == "Origine" for t in doc.tables)
    assert len(_rows(doc, "Atelier")) == 2  # but both versions stay in the history


# --- provenance that atelier 4 never receives (§12.3, §20.3) -------------------

def test_gap_provenance_keeps_framework_control_and_legal_split(repo):
    mid = repo.create_mission("Clinique Test", ["ISO27001", "RGPD"])
    repo.save_output(mid, 1, {"baseline_gaps_full": [
        {"gap_id": "BG-aaaa1111", "weakness": "Pas de MFA", "risk_categories": ["credential_access"],
         "controls": [{"framework": "ISO27001", "control_id": "A.5.17"}, {"framework": "RGPD", "control_id": "Art32"}]},
        {"gap_id": "BG-bbbb2222", "weakness": "Pas de registre", "risk_categories": [],
         "controls": [{"framework": "RGPD", "control_id": "Art30"}]},
    ], "unverified_controls": [{"control_id": "A.8.2", "framework": "ISO27001", "reason": "no_information"}]})

    doc = _annex(repo, mid)

    assert _rows(doc, "Écart") == [
        ["BG-aaaa1111", "Pas de MFA", "ISO27001", "A.5.17", "Sécurité"],
        ["BG-aaaa1111", "Pas de MFA", "RGPD", "Art32", "Sécurité"],
        ["BG-bbbb2222", "Pas de registre", "RGPD", "Art30", "Impact légal"],
    ]
    assert _rows(doc, "Référentiel", columns=3) == [["ISO27001", "A.8.2", "no_information"]]


# --- framework versions ---------------------------------------------------------

def test_framework_versions_default_to_the_plugin_manifests(repo):
    mid = repo.create_mission("Clinique Test", [])
    repo.save_output(mid, 0, {"applicable_frameworks": ["RGPD", "Referentiel-Inconnu"]})
    rgpd = next(p for p in discover_frameworks() if p.id.casefold() == "rgpd")

    data = collect_audit_annex_data(repo, mid, generated_at=FIXED_DATE)

    assert data.framework_versions == {"RGPD": rgpd.framework_version, "Referentiel-Inconnu": UNKNOWN_VERSION}
    assert rgpd.framework_version  # the manifest does carry a version


# --- edges ----------------------------------------------------------------------

def test_a_mission_with_no_work_still_produces_a_readable_annex(repo):
    mid = repo.create_mission("Mission vide", [])

    doc = _annex(repo, mid)

    lines = [p.text for p in doc.paragraphs]
    assert "Annexe d'audit" in lines
    assert "Mission vide" in lines
    assert "Aucune décision enregistrée." in lines
    assert "Aucun événement anormal." in lines


def test_unknown_mission_is_refused(repo):
    with pytest.raises(ValueError):
        collect_audit_annex_data(repo, "n-existe-pas")


def test_the_annex_is_written_to_a_real_word_file(repo, tmp_path):
    mid = repo.create_mission("Clinique Test", [])
    repo.save_output(mid, 1, {})
    data = collect_audit_annex_data(repo, mid, framework_versions={}, generated_at=FIXED_DATE)

    path = write_audit_annex(data, tmp_path / "sous-dossier" / "annexe.docx")

    assert path.exists()
    reopened = Document(str(path))
    assert "Clinique Test" in [p.text for p in reopened.paragraphs]
    assert ["Document généré le", FIXED_DATE] in _rows(reopened, "Rubrique")


# --- the house design (same cover and look as the progress reports) -------------

def test_the_cover_carries_the_annex_title_the_mission_and_the_month(repo):
    mid = repo.create_mission("Clinique Test", [])

    lines = [p.text for p in _annex(repo, mid).paragraphs]

    assert "ANNEXE D'AUDIT  ·  DOSSIER DE TRAÇABILITÉ" in lines
    assert "Annexe d'audit" in lines
    assert "Clinique Test" in lines
    assert "Projet OGSBC  ·  Annexe d'audit  ·  Octobre 2026" in lines
    assert not any("Atelier 1" == line or "Juillet" in line for line in lines)  # nothing left from the model report


def test_the_cover_names_nobody(repo):
    mid = repo.create_mission("Clinique Test", [])

    text = "\n".join(p.text for p in _annex(repo, mid).paragraphs)

    for name in ("ENCADRANT", "ÉQUIPE PROJET", "Jamal", "SAAD", "AGOUJIM", "Alae", "Salma", "Sanaa"):
        assert name not in text


def test_every_table_header_is_in_the_house_red(repo):
    mid = repo.create_mission("Clinique Test", [])

    doc = _annex(repo, mid)

    shades = {t.rows[0].cells[0]._tc.xpath("string(.//w:shd/@w:fill)") for t in doc.tables if len(t.columns) > 1}
    assert shades == {"BE1E2D"}


@pytest.mark.parametrize("iso, expected", [
    ("2026-10-10T16:00:00+00:00", "Octobre 2026"),
    ("2026-02-01T00:00:00+00:00", "Février 2026"),
    ("2027-08-31T23:59:59+00:00", "Août 2027"),
])
def test_french_month_and_year(iso, expected):
    assert french_month_year(iso) == expected
