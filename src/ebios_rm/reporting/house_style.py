"""House style of the project's Word reports — the design of the progress reports sent to the supervisor.

Both documents of the reporting agent (audit annex and mission report, conception
§20.1) open the same cover template, then add their content with the helpers
below, so they look like every report the team already sends.

The cover lives in templates/gabarit_rapport.docx; its five editable lines are
swapped by position (see ``Cover``). Everything after the cover is built here.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

TEMPLATE_PATH = Path(__file__).parent / "templates" / "gabarit_rapport.docx"

RED, GOLD = "BE1E2D", "CC9900"
BODY, LABEL = "34373B", "8F9BB1"
ALT_ROW, PINK, CREAM, WHITE = "F2F2F2", "FAEAEA", "FEF8E8", "FFFFFF"
GREEN, ORANGE = "1E7B3C", "B15C00"
BODY_FONT, TITLE_FONT = "Verdana", "Arial Black"

_MONTHS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
           "septembre", "octobre", "novembre", "décembre"]

# Positions of the editable lines in the template's body.
_COVER_KICKER, _COVER_TITLE, _COVER_SUBTITLE, _COVER_TAGLINE, _COVER_FOOTER = 2, 3, 4, 8, 10


@dataclass(frozen=True)
class Cover:
    kicker: str
    title: str
    subtitle: str
    tagline: str
    footer: str


def french_month_year(iso_date: str) -> str:
    """'2026-10-10T…' -> 'Octobre 2026' (no locale needed)."""
    moment = datetime.fromisoformat(iso_date)
    return f"{_MONTHS[moment.month - 1].capitalize()} {moment.year}"


def _replace_text(paragraph_element, text: str) -> None:
    runs = paragraph_element.findall(qn("w:r"))
    if not runs:
        return
    first = runs[0]
    for old in first.findall(qn("w:t")):
        first.remove(old)
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    first.append(node)
    for extra in runs[1:]:
        paragraph_element.remove(extra)


def _style_run(run, *, size=9.5, color=BODY, bold=False, italic=False, font=BODY_FONT) -> None:
    run.font.name = font
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold, run.italic = bold, italic


def _shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    for old in properties.findall(qn("w:shd")):
        properties.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _borders(cell, color="BBBBBB", size=4) -> None:
    properties = cell._tc.get_or_add_tcPr()
    for old in properties.findall(qn("w:tcBorders")):
        properties.remove(old)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:color"), color)
        borders.append(edge)
    properties.append(borders)


def _write_cell(cell, text: str, *, bold=False, color=BODY, size=9.5, align=None, shading=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    _style_run(paragraph.add_run(text), size=size, color=color, bold=bold)
    if shading:
        _shade(cell, shading)
    _borders(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


class HouseDocument:
    """A Word document that starts on the house cover and is filled with house-styled blocks."""

    def __init__(self, cover: Cover) -> None:
        self.document: DocxDocument = Document(str(TEMPLATE_PATH))
        body = self.document.element.body
        elements = list(body.iterchildren())
        for position, text in ((_COVER_KICKER, cover.kicker), (_COVER_TITLE, cover.title),
                               (_COVER_SUBTITLE, cover.subtitle), (_COVER_TAGLINE, cover.tagline),
                               (_COVER_FOOTER, cover.footer)):
            _replace_text(elements[position], text)

    # --- text ---

    def label(self, text: str, color: str = LABEL) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(14), Pt(2)
        _style_run(paragraph.add_run(text), size=6.5, color=color, bold=True, font=TITLE_FONT)

    def heading(self, text: str, *, level: int = 1, size: float | None = None) -> None:
        before, after, default = {1: (20, 10, 14), 2: (14, 7, 11)}.get(level, (10, 5, 10))
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(before), Pt(after)
        _style_run(paragraph.add_run(text), size=size or default, color=RED, bold=True, font=TITLE_FONT)

    def paragraph(self, text: str, *, italic: bool = False, color: str = BODY) -> None:
        block = self.document.add_paragraph()
        block.paragraph_format.space_before, block.paragraph_format.space_after = Pt(4), Pt(4)
        block.paragraph_format.line_spacing = 1.35
        block.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _style_run(block.add_run(text), size=9.5, color=color, italic=italic)

    def note(self, text: str) -> None:
        self.paragraph(text, italic=True, color=LABEL)

    def statement(self, text: str, number: int) -> None:
        """A sentence of the report, followed by a small reference number (the sources are listed at the end)."""
        block = self.document.add_paragraph()
        block.paragraph_format.space_before, block.paragraph_format.space_after = Pt(2), Pt(5)
        block.paragraph_format.line_spacing = 1.35
        block.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _style_run(block.add_run(text), size=9.5)
        marker = block.add_run(f" [{number}]")
        _style_run(marker, size=7, color=GOLD, bold=True)
        marker.font.superscript = True

    def reference_line(self, number: int, text: str) -> None:
        block = self.document.add_paragraph()
        block.paragraph_format.space_before, block.paragraph_format.space_after = Pt(0), Pt(3)
        _style_run(block.add_run(f"[{number}]  "), size=7.5, color=GOLD, bold=True)
        _style_run(block.add_run(text), size=7.5, color=LABEL)

    def spacer(self, points: int = 6) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(0), Pt(points)

    def page_break(self) -> None:
        run = self.document.add_paragraph().add_run()
        brk = OxmlElement("w:br")
        brk.set(qn("w:type"), "page")
        run._element.append(brk)

    def picture(self, png: bytes, width_inches: float = 6.4) -> None:
        from io import BytesIO  # noqa: PLC0415

        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(4), Pt(6)
        paragraph.add_run().add_picture(BytesIO(png), width=Inches(width_inches))

    def toc_line(self, number: str, title: str) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(6), Pt(2)
        _style_run(paragraph.add_run(number), size=11, color=RED, bold=True, font=TITLE_FONT)
        _style_run(paragraph.add_run("   " + title), size=11, bold=True, font=TITLE_FONT)

    # --- blocks ---

    def table(self, header: list[str], rows: list[list], widths: list[float] | None = None, *, size: float = 9.5) -> None:
        """Rows hold plain strings, or dicts {text, bold, color, align} for a styled cell."""
        table = self.document.add_table(rows=1 + len(rows), cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if widths:
            total = Inches(6.5).emu
            for index, percent in enumerate(widths):
                for row in table.rows:
                    row.cells[index].width = Emu(int(total * percent / 100))
        for cell, text in zip(table.rows[0].cells, header):
            _write_cell(cell, text, bold=True, color=WHITE, shading=RED, size=size)
        for r, values in enumerate(rows):
            shading = ALT_ROW if r % 2 == 1 else None
            for cell, value in zip(table.rows[r + 1].cells, values):
                style = value if isinstance(value, dict) else {"text": value}
                _write_cell(cell, style["text"], bold=style.get("bold", False), color=style.get("color", BODY),
                            align=style.get("align"), shading=shading, size=size)
        self.spacer(4)

    def callout(self, title: str, text: str, *, background: str = PINK, border: str = RED) -> None:
        table = self.document.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        cell.text = ""
        _style_run(cell.paragraphs[0].add_run(title), size=10, color=border, bold=True, font=TITLE_FONT)
        body = cell.add_paragraph()
        body.paragraph_format.space_before = Pt(4)
        _style_run(body.add_run(text), size=9.5)
        _shade(cell, background)
        properties = cell._tc.get_or_add_tcPr()
        for old in properties.findall(qn("w:tcBorders")):
            properties.remove(old)
        borders = OxmlElement("w:tcBorders")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "36")
        left.set(qn("w:color"), border)
        borders.append(left)
        for side in ("top", "bottom", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "4")
            edge.set(qn("w:color"), background)
            borders.append(edge)
        properties.append(borders)
        margins = OxmlElement("w:tcMar")
        for side, value in (("top", "160"), ("bottom", "160"), ("left", "240"), ("right", "240")):
            margin = OxmlElement(f"w:{side}")
            margin.set(qn("w:w"), value)
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        properties.append(margins)
        self.spacer(6)

    def save(self, path: str | Path) -> Path:
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(target))
        return target
